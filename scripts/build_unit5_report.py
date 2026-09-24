from __future__ import annotations

from pathlib import Path
import re
import subprocess
import textwrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from pygments import highlight
from pygments.lexers import JavascriptLexer
from pygments.formatters.img import ImageFormatter

ROOT = Path(__file__).resolve().parents[1]
OUTDIR = ROOT / "build" / "unit5-report"
OUTDIR.mkdir(parents=True, exist_ok=True)
DOCX = OUTDIR / "MSIT_5910_Unit5_Core_Logic_Testing_Version_Control.docx"

REPO_URL = "https://github.com/ctucker-web/ai-governance-capstone"
PR_URL = "https://github.com/ctucker-web/ai-governance-capstone/pull/3"
REL4_URL = "https://github.com/ctucker-web/ai-governance-capstone/releases/tag/v0.4.0-initial-mvp"
REL5_URL = "https://github.com/ctucker-web/ai-governance-capstone/releases/tag/v0.5.0-unit5"
CI_URL = "https://github.com/ctucker-web/ai-governance-capstone/actions/runs/36045015348"


def run(cmd: list[str]) -> str:
    p = subprocess.run(cmd, cwd=ROOT, text=True, capture_output=True, check=True)
    return re.sub(r"\x1b\[[0-9;]*m", "", p.stdout)


def screenshot_code(source: Path, start: int, end: int, title: str, target: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()[start - 1:end]
    numbered = "\n".join(lines)
    formatter = ImageFormatter(
        font_name="DejaVu Sans Mono",
        font_size=15,
        line_numbers=True,
        style="monokai",
        image_pad=20,
        line_pad=3,
        line_number_bg="#262626",
        line_number_fg="#A0A0A0",
    )
    target.write_bytes(highlight(f"// {title}\n{numbered}\n", JavascriptLexer(), formatter))


def screenshot_terminal(text: str, title: str, target: Path, max_lines: int = 32) -> None:
    text = re.sub(r"\x1b\[[0-9;]*m", "", text)
    raw = [x.rstrip() for x in text.splitlines() if x.strip()]
    lines: list[str] = []
    for line in raw:
        if len(lines) >= max_lines:
            break
        lines.extend(textwrap.wrap(line, 105) or [""])
    lines = lines[:max_lines]

    mono = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"
    mono_bold = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf"
    font = ImageFont.truetype(mono, 20)
    title_font = ImageFont.truetype(mono_bold, 23)
    width, line_h, header, margin = 1500, 30, 60, 32
    height = header + margin * 2 + line_h * len(lines)
    img = Image.new("RGB", (width, height), "#0d1117")
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, width, header), fill="#161b22")
    for x, c in [(22, "#ff5f56"), (47, "#ffbd2e"), (72, "#27c93f")]:
        d.ellipse((x, 21, x + 14, 35), fill=c)
    d.text((105, 14), title, font=title_font, fill="#f0f6fc")
    y = header + margin
    for line in lines:
        color = "#c9d1d9"
        if line.startswith("✔") or line.startswith("ok ") or "# pass 13" in line or '"tier": "High"' in line or '"decision": "Conditional"' in line:
            color = "#7ee787"
        elif line.startswith("==="):
            color = "#58a6ff"
        d.text((margin, y), line, font=font, fill=color)
        y += line_h
    img.save(target)


def screenshot_version_control(target: Path) -> None:
    sans = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    bold = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
    img = Image.new("RGB", (1500, 650), "white")
    d = ImageDraw.Draw(img)
    title = ImageFont.truetype(bold, 36)
    head = ImageFont.truetype(bold, 24)
    body = ImageFont.truetype(sans, 21)
    small = ImageFont.truetype(sans, 18)
    d.text((55, 40), "Version-control milestones", font=title, fill="#102a43")
    boxes = [
        ("main", "stable", 60),
        ("development", "integration", 390),
        ("feature/initial-mvp", "Unit 4", 720),
        ("feature/unit5-core-logic", "Unit 5", 1050),
    ]
    for label, sub, x in boxes:
        d.rounded_rectangle((x, 135, x + 280, 245), radius=15, fill="#f5f7fb", outline="#d0d5dd", width=3)
        d.text((x + 16, 155), label, font=head, fill="#1f2937")
        d.text((x + 16, 202), sub, font=small, fill="#667085")
    for x in (340, 670, 1000):
        d.line((x, 190, x + 44, 190), fill="#2563eb", width=5)
        d.polygon([(x + 44, 190), (x + 30, 181), (x + 30, 199)], fill="#2563eb")
    d.text((60, 310), "Tags / GitHub Releases", font=head, fill="#102a43")
    rows = [
        ("v0.4.0-initial-mvp", "Unit 4 - Initial Working MVP"),
        ("v0.5.0-unit5", "Unit 5 - Core Logic and Unit Testing"),
    ]
    y = 365
    for tag, name in rows:
        d.rounded_rectangle((60, y, 1440, y + 82), radius=14, fill="#eef4ff", outline="#b7ccf3", width=2)
        d.text((85, y + 20), tag, font=head, fill="#2563eb")
        d.text((475, y + 23), name, font=body, fill="#1f2937")
        y += 105
    d.text((60, 595), "Public repository: github.com/ctucker-web/ai-governance-capstone", font=small, fill="#667085")
    img.save(target)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(10)
    return p


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0.5)
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14 if level == 1 else 12)
    return p


def set_doc_style(doc: Document):
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


# Produce execution evidence from the actual checked-out branch.
test_output = run(["node", "--test", "--test-reporter=spec", "tests/riskEngine.test.js", "tests/workflow.test.js"])
demo_output = run(["node", "scripts/demo-core-logic.js"])

risk_img = OUTDIR / "figure1-risk-engine.png"
workflow_img = OUTDIR / "figure3-workflow.png"
test_img = OUTDIR / "figure4-tests.png"
demo_img = OUTDIR / "figure2-output.png"
vc_img = OUTDIR / "figure5-version-control.png"

screenshot_code(ROOT / "src/modules/risk/riskEngine.js", 40, 103, "Advisory scoring and mandatory escalation", risk_img)
screenshot_code(ROOT / "src/modules/workflow/workflow.js", 1, 49, "Human decision safeguards", workflow_img)

# Keep the terminal evidence focused on the actual results used in the report.
demo_focus = "\n".join(
    line for line in demo_output.splitlines()
    if line.startswith("===") or any(x in line for x in [
        '"score"', '"tier"', '"Sensitive or regulated data"',
        '"High impact on clients or employees"', '"advisoryOnly"',
        '"decision"', '"rationale"', '"action"', '"newStatus"', '"timestamp"'
    ])
)
screenshot_terminal(demo_focus, "npm run demo:logic - validated core output", demo_img, max_lines=26)
screenshot_terminal(test_output, "npm run test:verbose - 13 automated unit tests", test_img, max_lines=25)
screenshot_version_control(vc_img)

# Build report.
doc = Document()
set_doc_style(doc)

# Title page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(90)
r = p.add_run("Core Logic Implementation, Unit Testing, and Version Control Management")
r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(16)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("AI Governance Risk and Approval Platform for Nonprofit Organizations")
r.bold = True; r.font.size = Pt(14)
for line in ["Christopher Tucker", "MSIT 5910 Capstone Project", "University of the People", "September 24, 2026"]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(line)
doc.add_page_break()

add_heading(doc, "Part A: Core Logic Implementation and Output Validation")
add_heading(doc, "Alignment with Project Objectives", 2)
add_body(doc, "The primary objective of my capstone is to provide nonprofit organizations with a simple, consistent, and auditable way to identify and govern artificial intelligence use. Unit 4 demonstrated the end-to-end workflow in a browser. For this unit, I separated the most important business rules into testable modules: an explainable advisory risk engine and a human decision workflow. This directly supports the project's main objectives of consistent risk classification, transparency, human accountability, and traceability. The system is intentionally designed so that the algorithm informs review but never becomes the final governance authority.")

add_body(doc, "The risk engine evaluates five MVP dimensions: data sensitivity, impact on clients or employees, AI autonomy, limited human oversight, and limited vendor assurance. Each input is rated from 0 through 3. The initial equal-weight score maps 0-4 to Low, 5-8 to Moderate, 9-12 to High, and 13-15 to Restricted. The implementation also applies mandatory escalation rules because a simple total can hide important combinations. For example, high autonomy combined with high individual impact forces Restricted review, while sensitive data combined with limited vendor assurance requires at least High review. The returned object preserves the score, tier, factor contributions, explanations, escalation rules, and algorithm version. This makes the result understandable rather than opaque.")
doc.add_picture(str(risk_img), width=Inches(6.25))
add_caption(doc, "Figure 1. Essential risk-engine code showing score thresholds and mandatory escalation rules.")

add_body(doc, "Output validation used a synthetic AI-assisted documentation scenario. The sample produced a score of 10 and a High advisory tier because sensitive information and high impact were present. The output also retained the contributing factors and explicitly marked the result as advisory only. The same demonstration then recorded a Conditional human decision with a required rationale and generated a timestamped audit event. These outputs confirm that the implementation separates algorithmic assessment from human authority while preserving evidence for later review.")
doc.add_picture(str(demo_img), width=Inches(6.25))
add_caption(doc, "Figure 2. Core logic output showing the advisory risk result, human decision, and audit event.")

add_body(doc, "A second implementation challenge involved enforcing governance rules outside the user interface. Hiding an approval button is not sufficient authorization. I therefore created workflow-level validation that only Reviewer or Administrator roles may record a final decision, prevents a requester from reviewing their own submission, restricts final decisions to Approved, Conditional, or Rejected, and requires a rationale. This makes the rule portable to a future server-side implementation. The main challenge was balancing MVP simplicity with meaningful safeguards. I resolved it by keeping the algorithm deterministic and small while adding explicit escalation and authorization rules instead of introducing machine learning or complex automation.")
doc.add_picture(str(workflow_img), width=Inches(6.25))
add_caption(doc, "Figure 3. Workflow safeguards enforcing reviewer authority, separation of duties, and rationale.")

add_heading(doc, "Part B: Unit Testing of Functional Modules")
add_heading(doc, "Testing Approach and Results", 2)
add_body(doc, "I used Node.js's built-in node:test framework with strict assertions. This is an appropriate JavaScript unit-testing framework and avoids adding unnecessary dependencies to the early MVP. I performed both black-box and white-box unit testing. Black-box tests validate observable behavior without relying on internal structure, such as verifying that a score of 5 returns Moderate risk or that a Requester cannot record a decision. White-box tests deliberately exercise known internal branches, including mandatory escalation rules and the separation-of-duties condition that prevents self-approval. Boundary and negative testing were also represented by checking threshold values and rejecting factor inputs outside the permitted 0-3 range.")

add_body(doc, "The automated suite contains 13 tests. Seven focus on the risk engine: Low behavior, the Moderate and High boundaries, Restricted escalation, vendor-assurance escalation, retained explanations, and invalid input. Six focus on workflow and audit logic: Requester denial, Reviewer permission, self-review prevention, required rationale, decision record creation, and audit-event creation. The local execution and the GitHub Actions CI run both completed with 13 tests passed and zero failures. The same CI workflow also runs the repeatable core-logic demonstration, so a code change that breaks the algorithm can be detected before integration.")
doc.add_picture(str(test_img), width=Inches(6.25))
add_caption(doc, "Figure 4. Unit-test execution showing 13 tests passed and zero failures.")

add_body(doc, "Testing identified several issues that were easy to miss in the visual prototype. First, UI-only role restrictions did not protect the underlying business rule, so authorization validation was moved into the workflow module. Second, reviewer rationale needed to be programmatically required rather than treated as an optional form behavior. Third, threshold tests showed why score-only classification was insufficient for high-consequence combinations, which reinforced the mandatory escalation design. These changes demonstrate how unit tests improve reliability: defects are isolated close to the module where they originate, and each module can change without requiring the entire application to be retested manually. Black-box tests protect expected behavior, white-box tests cover important branches, boundary tests protect thresholds, negative tests verify rejection of invalid actions or inputs, and regression tests prevent previously corrected behavior from returning.")

add_heading(doc, "Part C: Version Control Management")
add_heading(doc, "Repository, Branching, Tags, and Releases", 2)
p = doc.add_paragraph(); p.paragraph_format.line_spacing = 2
p.add_run("Public GitHub repository: ")
add_hyperlink(p, REPO_URL, REPO_URL)

add_body(doc, "The repository uses a simple layered branching model. main represents the stable project baseline, development is the integration branch, and short-lived feature branches isolate significant work. The Unit 4 prototype was developed on feature/initial-mvp. Unit 5 was developed on feature/unit5-core-logic from the Unit 4 milestone so that algorithm and test changes remained separate and reviewable. Pull Request #3 contains the Unit 5 change set. This approach provides traceability from a requirement to a feature branch, commit, automated test run, pull request, tag, and release.")
doc.add_picture(str(vc_img), width=Inches(6.25))
add_caption(doc, "Figure 5. Branch and milestone structure used for the Unit 4 and Unit 5 work.")

add_body(doc, "I created two milestone tags and corresponding GitHub releases. v0.4.0-initial-mvp records the first working browser MVP. v0.5.0-unit5 records the implementation of modular core logic and automated unit testing. The releases were created through a GitHub Actions milestone workflow, demonstrating an advanced version-control practice in which repeatable automation manages release metadata instead of relying only on manual steps. The Unit 5 CI run also completed successfully after checking required documentation, executing the unit tests, and validating the example output.")

p = doc.add_paragraph(); p.paragraph_format.line_spacing = 2
p.add_run("Unit 5 Pull Request: "); add_hyperlink(p, "Pull Request #3", PR_URL)
p.add_run("\nUnit 4 release: "); add_hyperlink(p, "v0.4.0-initial-mvp", REL4_URL)
p.add_run("\nUnit 5 release: "); add_hyperlink(p, "v0.5.0-unit5", REL5_URL)
p.add_run("\nSuccessful CI run: "); add_hyperlink(p, "GitHub Actions run 10", CI_URL)

add_body(doc, "Tags provide fixed references to milestone commits, making it possible to reproduce or inspect the exact state submitted at a particular stage. Releases add human-readable milestone notes and downloadable source archives. Feature branches reduce the risk of destabilizing integration work, while pull requests create a visible review boundary. Together, branches, CI, tags, and releases make progress easier to monitor and provide evidence that the capstone is being developed as a controlled software-engineering project rather than as a collection of untracked files.")

add_heading(doc, "Conclusion")
add_body(doc, "Unit 5 strengthened the MVP by moving central governance rules into modular, tested code. The advisory risk engine remains explainable and subordinate to human judgment, while workflow validation protects reviewer accountability and auditability. Thirteen automated tests and a successful GitHub Actions run provide repeatable evidence that the modules behave as expected. The version-control milestones also establish a clear history from initial MVP through tested core logic, creating a stronger foundation for later database, authentication, integration, and evaluation work.")

# Ensure all body text is Times New Roman 12.
for p in doc.paragraphs:
    for run_ in p.runs:
        if not run_.font.name:
            run_.font.name = "Times New Roman"
        if not run_.font.size:
            run_.font.size = Pt(12)

# Basic word-count diagnostic for the assignment narrative only.
text = "\n".join(p.text for p in doc.paragraphs[5:])
words = re.findall(r"\b[\w'-]+\b", text)
print(f"Approximate report word count including captions/links: {len(words)}")
print(f"Tests passed evidence captured: {'pass 13' in test_output or 'tests 13' in test_output}")

doc.save(DOCX)
print(DOCX)
